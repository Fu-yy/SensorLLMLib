import os
import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Border, Side


# ================= 辅助函数 =================

def get_stats(values):
    """
    计算均值和标准差。
    注意：这里返回的是 std，不是真正的 95% CI。
    """
    if not values:
        return np.nan, np.nan

    a = np.array(values, dtype=float)
    n = len(a)
    mean = np.mean(a)

    if n < 2:
        return mean, 0.0

    std = np.std(a, ddof=1)
    return mean, std


def extract_metrics_from_line(line):
    """
    提取数值指标，例如：
    loss:0.123 acc:0.987 f1_macro:0.901
    支持小数、负数、科学计数法。
    """
    pattern = r"([a-zA-Z0-9_]+):\s*(-?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)"
    matches = re.findall(pattern, line)

    metrics = {}
    for key, value in matches:
        metrics[key] = float(value)

    return metrics


def extract_model_name_from_line(line):
    """
    从日志中提取模型名称。
    例如：
    Model: xxx_model
    """
    pattern = r"Model:\s+([a-zA-Z0-9_\-\.]+)"
    match = re.search(pattern, line)

    if match:
        return match.group(1)

    return None


def remove_formatting(excel_path):
    """清除 Excel 默认样式"""
    try:
        wb = load_workbook(excel_path)
        ws = wb.active

        no_bold = Font(bold=False)
        no_border = Border(
            left=Side(style=None),
            right=Side(style=None),
            top=Side(style=None),
            bottom=Side(style=None)
        )

        for row in ws.iter_rows():
            for cell in row:
                cell.font = no_bold
                cell.border = no_border

        wb.save(excel_path)

    except Exception as e:
        print(f"清除样式忽略: {e}")


# ================= 单个 batch 的解析 =================

def collect_one_batch(batch_folder):
    """
    解析一个 batch 文件夹。

    输入：
        batch_folder:
            run_log/batch_xxx

    输出：
        row_id:
            batch名 + 模型名

        row_data:
            {
                (dataset, metric, stat): value
            }

        dataset_list:
            当前 batch 中检测到的数据集

        metric_order:
            当前 batch 中检测到的指标顺序
    """

    batch_base_name = os.path.basename(os.path.normpath(batch_folder))
    target_prefix = "[Stage2-Test]"

    print(f"\n开始解析 batch: {batch_base_name}")
    print(f"路径: {batch_folder}")

    data_store = {}
    detected_model_name = None

    for root, dirs, files in os.walk(batch_folder):
        if "stage2.log" not in files:
            continue

        dataset_name = os.path.basename(root)
        log_path = os.path.join(root, "stage2.log")

        if dataset_name not in data_store:
            data_store[dataset_name] = {}

        try:
            with open(log_path, "r", encoding="utf-8") as f:
                for line in f:
                    # 1. 提取模型名
                    if detected_model_name is None:
                        model_str = extract_model_name_from_line(line)
                        if model_str:
                            detected_model_name = model_str

                    # 2. 提取 Stage2-Test 测试指标
                    if target_prefix in line:
                        content = line.split(target_prefix, 1)[1]
                        metrics = extract_metrics_from_line(content)

                        for k, v in metrics.items():
                            if k not in data_store[dataset_name]:
                                data_store[dataset_name][k] = []
                            data_store[dataset_name][k].append(v)

        except Exception as e:
            print(f"读取出错: {log_path} -> {e}")

    if not data_store:
        print(f"警告：{batch_base_name} 未找到有效 Stage2-Test 数据")
        return None, None, [], []

    # 生成行 ID
    if detected_model_name:
        row_id = f"{batch_base_name}_{detected_model_name}"
    else:
        row_id = batch_base_name
        print(f"警告：{batch_base_name} 未找到 Model 字段，仅使用 batch 文件夹名作为 ID")

    # 收集所有指标
    all_metrics = set()
    for ds_metrics in data_store.values():
        all_metrics.update(ds_metrics.keys())

    preferred_order = [
        "loss",
        "acc",
        "f1_macro",
        "recall_macro",
        "precision_macro",
        "infer_total_time",
        "infer_ms_per_sample",
        "infer_samples_per_sec",
        "infer_total_samples"
    ]

    remaining_metrics = sorted([m for m in all_metrics if m not in preferred_order])
    metric_order = [m for m in preferred_order if m in all_metrics] + remaining_metrics

    dataset_list = sorted(data_store.keys())

    row_data = {}

    for dataset_name in dataset_list:
        metrics_data = data_store[dataset_name]

        for metric in metric_order:
            values = metrics_data.get(metric, [])
            mean, std = get_stats(values)

            row_data[(dataset_name, metric, "Mean")] = mean
            row_data[(dataset_name, metric, "Std")] = std

    print(f"完成解析: {row_id}")
    print(f"检测到数据集: {dataset_list}")

    return row_id, row_data, dataset_list, metric_order


# ================= 多个 batch 汇总到一个 Excel =================

def process_all_batches(total_root_folder, output_excel):
    """
    解析 run_log 总目录下的多个 batch 文件夹，并汇总到一个 Excel。

    目录示例：
        run_log/
            batch_20260505_xxx/
                mhealth/stage2.log
                pamap2/stage2.log
            batch_20260506_xxx/
                mhealth/stage2.log
                pamap2/stage2.log
    """

    print("=" * 80)
    print(f"总目录: {total_root_folder}")
    print(f"输出文件: {output_excel}")
    print("=" * 80)

    if not os.path.exists(total_root_folder):
        print(f"错误：总目录不存在：{total_root_folder}")
        return

    # 找到总目录下的所有 batch 文件夹
    batch_folders = []

    for name in sorted(os.listdir(total_root_folder)):
        path = os.path.join(total_root_folder, name)

        if not os.path.isdir(path):
            continue

        # 推荐只处理 batch_ 开头的文件夹，避免误扫其他文件
        if name.startswith("batch_"):
            batch_folders.append(path)

    if not batch_folders:
        print("未找到 batch_ 开头的文件夹。")
        return

    print(f"检测到 {len(batch_folders)} 个 batch 文件夹：")
    for p in batch_folders:
        print("  -", os.path.basename(p))

    all_rows = {}
    all_datasets = set()
    all_metrics = set()

    # 逐个解析 batch
    for batch_folder in batch_folders:
        row_id, row_data, dataset_list, metric_order = collect_one_batch(batch_folder)

        if row_id is None:
            continue

        all_rows[row_id] = row_data
        all_datasets.update(dataset_list)
        all_metrics.update(metric_order)

    if not all_rows:
        print("没有任何有效数据，跳过生成 Excel。")
        return

    # 统一所有 batch 的列顺序
    preferred_order = [
        "loss",
        "acc",
        "f1_macro",
        "recall_macro",
        "precision_macro",
        "infer_total_time",
        "infer_ms_per_sample",
        "infer_samples_per_sec",
        "infer_total_samples"
    ]

    remaining_metrics = sorted([m for m in all_metrics if m not in preferred_order])
    final_metric_order = [m for m in preferred_order if m in all_metrics] + remaining_metrics

    final_dataset_order = sorted(all_datasets)

    desired_columns = pd.MultiIndex.from_product(
        [
            final_dataset_order,
            final_metric_order,
            ["Mean", "Std"]
        ],
        names=["Dataset", "Metric", "Stat"]
    )

    # 生成 DataFrame
    df = pd.DataFrame.from_dict(all_rows, orient="index")
    df = df.reindex(columns=desired_columns)
    df.index.name = "ID"

    # 写入 Excel
    try:
        output_dir = os.path.dirname(output_excel)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        df.to_excel(output_excel)

        print("\n" + "=" * 80)
        print(f"所有实验数据已写入：{output_excel}")
        print(f"共汇总实验数量：{len(df)}")
        print(f"数据集数量：{len(final_dataset_order)}")
        print(f"指标数量：{len(final_metric_order)}")
        print("=" * 80)

        remove_formatting(output_excel)

    except Exception as e:
        print(f"保存 Excel 失败: {e}")


# ================= 主入口 =================

if __name__ == "__main__":

    # 这里填你图片里的总文件夹，也就是 run_log
    TOTAL_ROOT_DIRECTORY = r"D:\fuy\MyCode\SensorLLMLib_v2_compare\run_log"
    # TOTAL_ROOT_DIRECTORY = r"/root/autodl-tmp/SensorLLMLib_v2/run_log"

    # 最终所有实验统一输出到这一个 Excel
    OUTPUT_FILE = r"D:\fuy\MyCode\SensorLLMLib_v2_compare\run_log\all_experiment_results_compare.xlsx"
    # OUTPUT_FILE = r"/root/autodl-tmp/SensorLLMLib_v2/run_log/all_experiment_results_compare.xlsx"

    process_all_batches(TOTAL_ROOT_DIRECTORY, OUTPUT_FILE)